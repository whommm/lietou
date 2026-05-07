"""Excel read/write service for candidate workflow."""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..models import CandidateExcelRecord


class CandidateExcelError(Exception):
    """Raised when Excel read/write fails."""


class CandidateExcelService:
    """Manage workbook creation, row updates, and match result writes."""

    SHEET_NAME = "候选人"
    EXPORT_DIR = os.path.join("exports", "candidates")
    CAPTURE_STATUS_PENDING = "待抓取"
    CAPTURE_STATUS_SUCCESS = "抓取成功"
    CAPTURE_STATUS_FAILED = "抓取失败"
    CAPTURE_STATUS_PARTIAL = "部分成功"
    MATCHABLE_CAPTURE_STATUSES = {CAPTURE_STATUS_SUCCESS, CAPTURE_STATUS_PARTIAL}
    HEADERS = [
        "序号",
        "姓名",
        "年龄",
        "来源关键词",
        "页码",
        "排名",
        "简历链接",
        "简历抓取状态",
        "简历详情",
        "匹配档位",
        "匹配详情",
        "抓取时间",
        "匹配时间",
        "打招呼状态",
        "打招呼时间",
        "打招呼消息",
        "人才标签",
        "联系方式",
    ]
    HEADER_ALIASES = {
        "匹配档位": ["匹配度分数"],
    }
    COLUMN_WIDTHS = {
        "A": 8,
        "B": 14,
        "C": 8,
        "D": 22,
        "E": 8,
        "F": 8,
        "G": 38,
        "H": 14,
        "I": 80,
        "J": 12,
        "K": 80,
        "L": 20,
        "M": 20,
        "N": 16,
        "O": 20,
        "P": 56,
        "Q": 18,
        "R": 36,
    }
    HEADER_INDEX = {name: index + 1 for index, name in enumerate(HEADERS)}

    def __init__(self, workspace_root: Optional[str] = None):
        self.workspace_root = workspace_root or os.getcwd()

    @staticmethod
    def now_text() -> str:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    @staticmethod
    def build_filename(job_name: str) -> str:
        normalized = re.sub(r'[\\/:*?"<>|]+', "_", (job_name or "候选人").strip())
        normalized = normalized.strip(" .") or "候选人"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return "{}_{}.xlsx".format(normalized, timestamp)

    def build_export_path(self, job_name: str) -> str:
        export_dir = os.path.join(self.workspace_root, self.EXPORT_DIR)
        os.makedirs(export_dir, exist_ok=True)
        return os.path.join(export_dir, self.build_filename(job_name))

    def create_workbook(self, job_name: str) -> str:
        path = self.build_export_path(job_name)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.SHEET_NAME
        sheet.append(self.HEADERS)
        self._apply_layout(sheet)
        workbook.save(path)
        workbook.close()
        return path

    def append_candidate_row(self, file_path: str, row_data: Dict[str, object]) -> int:
        try:
            workbook = load_workbook(file_path)
            sheet = workbook[self.SHEET_NAME]
            self._ensure_schema(sheet)
            row_index = sheet.max_row + 1
            values = [self._get_row_value(row_data, header) for header in self.HEADERS]
            sheet.append(values)
            self._apply_row_style(sheet, row_index)
            workbook.save(file_path)
            workbook.close()
            return row_index
        except PermissionError as exc:
            raise CandidateExcelError(
                "Excel 文件被占用，请关闭后重试：{}".format(file_path)
            ) from exc
        except Exception as exc:
            raise CandidateExcelError(
                "写入 Excel 失败：{}".format(exc)
            ) from exc

    def update_candidate_detail(
        self,
        file_path: str,
        row_index: int,
        resume_text: str,
        capture_status: str,
    ) -> None:
        try:
            workbook = load_workbook(file_path)
            sheet = workbook[self.SHEET_NAME]
            self._ensure_schema(sheet)
            sheet.cell(row=row_index, column=self.HEADER_INDEX["简历详情"]).value = (
                resume_text or ""
            )
            sheet.cell(
                row=row_index, column=self.HEADER_INDEX["简历抓取状态"]
            ).value = capture_status
            contact_info = self.extract_contact_info(resume_text)
            if contact_info:
                sheet.cell(row=row_index, column=self.HEADER_INDEX["联系方式"]).value = (
                    contact_info
                )
            talent_tags = self.extract_talent_tags(resume_text)
            if talent_tags:
                existing = sheet.cell(row=row_index, column=self.HEADER_INDEX["人才标签"]).value
                sheet.cell(row=row_index, column=self.HEADER_INDEX["人才标签"]).value = (
                    self._merge_tags(str(existing or ""), talent_tags)
                )
            workbook.save(file_path)
            workbook.close()
        except PermissionError as exc:
            raise CandidateExcelError(
                "Excel 文件被占用，请关闭后重试：{}".format(file_path)
            ) from exc
        except Exception as exc:
            raise CandidateExcelError(
                "更新 Excel 失败：{}".format(exc)
            ) from exc

    def write_match_result(
        self,
        file_path: str,
        row_index: int,
        tier: Optional[str],
        detail: str,
        matched_at: Optional[str] = None,
    ) -> None:
        try:
            workbook = load_workbook(file_path)
            sheet = workbook[self.SHEET_NAME]
            self._ensure_schema(sheet)
            sheet.cell(row=row_index, column=self.HEADER_INDEX["匹配档位"]).value = tier
            sheet.cell(row=row_index, column=self.HEADER_INDEX["匹配详情"]).value = (
                detail or ""
            )
            sheet.cell(row=row_index, column=self.HEADER_INDEX["匹配时间"]).value = (
                matched_at or self.now_text()
            )
            workbook.save(file_path)
            workbook.close()
        except PermissionError as exc:
            raise CandidateExcelError(
                "Excel 文件被占用，请关闭后重试：{}".format(file_path)
            ) from exc
        except Exception as exc:
            raise CandidateExcelError(
                "回写匹配结果到 Excel 失败：{}".format(exc)
            ) from exc

    def write_greeting_result(
        self,
        file_path: str,
        row_index: int,
        status: str,
        message: str = "",
        greeted_at: Optional[str] = None,
    ) -> None:
        """Write the greeting result back to the candidate row."""
        try:
            workbook = load_workbook(file_path)
            sheet = workbook[self.SHEET_NAME]
            self._ensure_schema(sheet)
            sheet.cell(row=row_index, column=self.HEADER_INDEX["打招呼状态"]).value = (
                status or ""
            )
            sheet.cell(row=row_index, column=self.HEADER_INDEX["打招呼时间"]).value = (
                greeted_at or self.now_text()
            )
            sheet.cell(row=row_index, column=self.HEADER_INDEX["打招呼消息"]).value = (
                message or ""
            )
            workbook.save(file_path)
            workbook.close()
        except PermissionError as exc:
            raise CandidateExcelError(
                "Excel 文件被占用，请关闭后重试：{}".format(file_path)
            ) from exc
        except Exception as exc:
            raise CandidateExcelError(
                "回写打招呼结果到 Excel 失败：{}".format(exc)
            ) from exc

    def load_candidates(self, file_path: str) -> List[CandidateExcelRecord]:
        workbook = load_workbook(file_path)
        sheet = workbook[self.SHEET_NAME]
        self._ensure_schema(sheet)
        records = []
        for row_index in range(2, sheet.max_row + 1):
            values = self._row_to_dict(sheet, row_index)
            if not any(
                str(values.get(header) or "").strip() for header in self.HEADERS
            ):
                continue
            records.append(self._build_record(row_index, values))
        workbook.close()
        return records

    def load_matchable_candidates(self, file_path: str) -> List[CandidateExcelRecord]:
        records = []
        for record in self.load_candidates(file_path):
            if (
                record.resume_text or ""
            ).strip() and record.capture_status in self.MATCHABLE_CAPTURE_STATUSES:
                records.append(record)
        return records

    def load_matchable_candidates_by_rows(
        self, file_path: str, row_indexes: List[int]
    ) -> List[CandidateExcelRecord]:
        """Load matchable candidates whose Excel rows are in ``row_indexes``."""
        target_rows = {int(row) for row in (row_indexes or []) if row}
        if not target_rows:
            return []
        return [
            record
            for record in self.load_matchable_candidates(file_path)
            if record.row_index in target_rows
        ]

    def load_matchable_candidates_by_source_keyword(
        self, file_path: str, source_keyword: str
    ) -> List[CandidateExcelRecord]:
        """Load matchable candidates captured from one search query."""
        source_keyword = (source_keyword or "").strip()
        if not source_keyword:
            return []
        return [
            record
            for record in self.load_matchable_candidates(file_path)
            if (record.source_keyword or "").strip() == source_keyword
        ]

    def count_matchable_candidates(self, file_path: str) -> int:
        return len(self.load_matchable_candidates(file_path))

    def load_greetable_candidates(
        self,
        file_path: str,
        tiers: Optional[set] = None,
        require_gold_collar: bool = False,
        require_no_contact: bool = False,
    ) -> List[CandidateExcelRecord]:
        """Load candidates with explicit high-priority match tiers for greeting."""
        target_tiers = {str(t).strip().upper() for t in (tiers or {"A", "B"})}
        records = []
        for record in self.load_candidates(file_path):
            tier = (record.match_tier or "").strip().upper()
            if (
                record.capture_status in self.MATCHABLE_CAPTURE_STATUSES
                and record.profile_url
                and tier in target_tiers
                and record.greeting_status not in ("发送成功", "已打过招呼")
                and (
                    not require_gold_collar
                    or self.is_gold_collar_candidate(record)
                )
                and (
                    not require_no_contact
                    or not self.has_contact_info(record)
                )
            ):
                records.append(record)
        return records

    def count_greetable_candidates(
        self,
        file_path: str,
        tiers: Optional[set] = None,
        require_gold_collar: bool = False,
        require_no_contact: bool = False,
    ) -> int:
        return len(
            self.load_greetable_candidates(
                file_path,
                tiers=tiers,
                require_gold_collar=require_gold_collar,
                require_no_contact=require_no_contact,
            )
        )

    def _build_record(
        self, row_index: int, values: Dict[str, object]
    ) -> CandidateExcelRecord:
        tier_value = values.get("匹配档位")
        match_tier = self._normalize_tier(tier_value)
        match_score = None
        if tier_value not in (None, ""):
            try:
                match_score = int(float(tier_value))
            except (TypeError, ValueError):
                match_score = None
        if not match_tier and match_score is not None:
            match_tier = self._tier_from_score(match_score)
        if not match_tier:
            match_tier = self._tier_from_detail(str(values.get("匹配详情") or ""))
        return CandidateExcelRecord(
            row_index=row_index,
            sequence=self._to_int(values.get("序号"), fallback=row_index - 1),
            name=str(values.get("姓名") or ""),
            age=str(values.get("年龄") or ""),
            source_keyword=str(values.get("来源关键词") or ""),
            page_number=self._to_int(values.get("页码")),
            rank_index=self._to_int(values.get("排名")),
            profile_url=str(values.get("简历链接") or ""),
            capture_status=str(values.get("简历抓取状态") or ""),
            resume_text=str(values.get("简历详情") or ""),
            match_tier=match_tier,
            match_score=match_score,
            match_detail=str(values.get("匹配详情") or ""),
            captured_at=str(values.get("抓取时间") or ""),
            matched_at=str(values.get("匹配时间") or ""),
            greeting_status=str(values.get("打招呼状态") or ""),
            greeted_at=str(values.get("打招呼时间") or ""),
            greeting_message=str(values.get("打招呼消息") or ""),
            talent_tags=str(values.get("人才标签") or ""),
            contact_info=str(values.get("联系方式") or ""),
        )

    @staticmethod
    def _to_int(value, fallback: Optional[int] = None) -> Optional[int]:
        if value in (None, ""):
            return fallback
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return fallback

    def _row_to_dict(self, sheet, row_index: int) -> Dict[str, object]:
        data = {}
        for header, column_index in self.HEADER_INDEX.items():
            data[header] = sheet.cell(row=row_index, column=column_index).value
        return data

    @classmethod
    def _get_row_value(cls, row_data: Dict[str, object], header: str):
        if header in row_data:
            return row_data.get(header, "")
        for alias in cls.HEADER_ALIASES.get(header, []):
            if alias in row_data:
                return row_data.get(alias, "")
        return ""

    def _ensure_schema(self, sheet) -> None:
        """Keep existing workbooks compatible with the current header names."""
        for header, column_index in self.HEADER_INDEX.items():
            current_value = sheet.cell(row=1, column=column_index).value
            aliases = self.HEADER_ALIASES.get(header, [])
            if current_value in (None, "", *aliases):
                sheet.cell(row=1, column=column_index).value = header
        self._apply_layout(sheet)

    @staticmethod
    def _normalize_tier(value) -> str:
        if value is None:
            return ""
        text = str(value).strip().upper()
        return text if text in ("A", "B", "C", "D") else ""

    @staticmethod
    def _tier_from_score(score: int) -> str:
        if score >= 80:
            return "A"
        if score >= 60:
            return "B"
        if score >= 40:
            return "C"
        return "D"

    @staticmethod
    def _tier_from_detail(detail: str) -> str:
        match = re.search(
            r"(?:档位判定|匹配档位|档位|等级|级别)\s*[:：]?\s*([ABCD])",
            detail or "",
            re.I,
        )
        return match.group(1).upper() if match else ""

    @staticmethod
    def extract_contact_info(text: str) -> str:
        """Extract concrete contact values from captured text."""
        text = text or ""
        contacts = []
        patterns = [
            r"1[3-9]\d{9}",
            r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
            r"(?:微信|WeChat|wechat|VX|vx)[：:\s]*([A-Za-z0-9_\-]{5,})",
            r"(?:QQ|qq)[：:\s]*(\d{5,})",
        ]
        for pattern in patterns:
            for match in re.finditer(pattern, text):
                value = match.group(1) if match.lastindex else match.group(0)
                if value and value not in contacts:
                    contacts.append(value)
        return " / ".join(contacts)

    @staticmethod
    def extract_talent_tags(text: str) -> str:
        """Extract known candidate tags that matter for downstream workflow."""
        text = text or ""
        tags = []
        for marker in ("金领", "金领人才", "金领简历"):
            if marker in text and "金领" not in tags:
                tags.append("金领")
        return " / ".join(tags)

    @classmethod
    def has_contact_info(cls, record: CandidateExcelRecord) -> bool:
        """Return True only when concrete contact values are available."""
        explicit = (record.contact_info or "").strip()
        no_contact_markers = {"", "无", "暂无", "未获取", "无法获取", "不能获取", "未知"}
        if explicit and explicit not in no_contact_markers:
            return True
        return bool(cls.extract_contact_info(record.resume_text))

    @classmethod
    def is_gold_collar_candidate(cls, record: CandidateExcelRecord) -> bool:
        """Detect Liepin gold-collar candidates from explicit tags or text."""
        text = "\n".join(
            [
                record.talent_tags or "",
                record.source_keyword or "",
                record.resume_text or "",
                record.match_detail or "",
            ]
        )
        return "金领" in text

    @staticmethod
    def _merge_tags(existing: str, new_tags: str) -> str:
        tags = []
        for item in re.split(r"[/／、,，;；\s]+", "{} {}".format(existing, new_tags)):
            value = item.strip()
            if value and value not in tags:
                tags.append(value)
        return " / ".join(tags)

    def _apply_layout(self, sheet) -> None:
        sheet.freeze_panes = "A2"
        header_fill = PatternFill("solid", fgColor="DCE7FF")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for column, width in self.COLUMN_WIDTHS.items():
            sheet.column_dimensions[column].width = width

    def _apply_row_style(self, sheet, row_index: int) -> None:
        wrap_columns = [
            self.HEADER_INDEX["简历详情"],
            self.HEADER_INDEX["匹配详情"],
            self.HEADER_INDEX["打招呼消息"],
        ]
        for column_index in wrap_columns:
            sheet.cell(row=row_index, column=column_index).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
